#!/usr/bin/env python3
"""
PPT向けインフォグラフィック生成エントリポイント。

使い方:
  python generate_infographic.py
  python generate_infographic.py -o output/my_deck.pptx
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation

from infographic.content_loader import build_from_json
from infographic.slides import (
    add_comparison_slide,
    add_kpi_slide,
    add_matrix_slide,
    add_pillars_slide,
    add_process_slide,
    add_takeaways_slide,
    add_timeline_slide,
    add_title_slide,
)
from infographic.theme import Theme


def build_sample_deck() -> Presentation:
    """サンプルの日本語インフォグラフィックデッキを生成。"""
    prs = Presentation()
    prs.slide_width = Theme.SLIDE_WIDTH
    prs.slide_height = Theme.SLIDE_HEIGHT

    add_title_slide(
        prs,
        title="業務改善プロジェクト\nインフォグラフィック概要",
        subtitle="現状課題 → 施策 → 効果 を一目で伝えるプレゼン資料",
        footer="Sample Infographic Deck  |  編集してご利用ください",
    )

    add_kpi_slide(
        prs,
        title="主要KPIサマリー",
        subtitle="施策導入後の定量効果（サンプル数値）",
        kpis=[
            {"value": "▲32%", "label": "作業時間削減", "note": "定型業務の自動化により"},
            {"value": "98%", "label": "顧客満足度", "note": "前年比 +12pt"},
            {"value": "1.8x", "label": "処理スピード", "note": "ボトルネック解消後"},
            {"value": "¥24M", "label": "年間コスト削減", "note": "人件費・手戻り含む"},
        ],
    )

    add_process_slide(
        prs,
        title="推進プロセス",
        subtitle="仮説検証を回しながら、確実に成果へつなげる5ステップ",
        steps=[
            {
                "title": "現状把握",
                "desc": "現場ヒアリングとデータ分析で、課題とボトルネックを可視化します。",
            },
            {
                "title": "論点整理",
                "desc": "優先度・影響度で論点を絞り、施策の仮説を立てます。",
            },
            {
                "title": "施策設計",
                "desc": "小さな実験単位に分解し、成功条件とKPIを定義します。",
            },
            {
                "title": "実行・検証",
                "desc": "パイロット導入後、週次で効果測定と改善を繰り返します。",
            },
            {
                "title": "横展開",
                "desc": "再現可能な仕組みとして標準化し、他部署へ展開します。",
            },
        ],
    )

    add_comparison_slide(
        prs,
        title="Before / After 比較",
        subtitle="業務プロセスの変化を対比して示す",
        left_panel={
            "title": "Before（改善前）",
            "items": [
                "属人的なノウハウに依存",
                "手作業の転記・確認が多い",
                "進捗が見えず遅延が発覚しづらい",
                "問い合わせ対応に平均2日",
                "改善活動が散発的",
            ],
            "color": Theme.ACCENT_4,
        },
        right_panel={
            "title": "After（改善後）",
            "items": [
                "標準手順とチェックリストを整備",
                "自動連携で転記ミスをゼロ化",
                "ダッシュボードでリアルタイム可視化",
                "一次回答を当日中に完了",
                "週次レビューで継続改善",
            ],
            "color": Theme.ACCENT_2,
        },
    )

    add_timeline_slide(
        prs,
        title="ロードマップ",
        subtitle="四半期単位のマイルストーン（サンプル）",
        milestones=[
            {
                "period": "Q1",
                "title": "現状診断",
                "desc": "課題の洗い出しとKPI定義",
            },
            {
                "period": "Q2",
                "title": "パイロット",
                "desc": "対象業務で小さく検証",
            },
            {
                "period": "Q3",
                "title": "本格導入",
                "desc": "運用定着と教育展開",
            },
            {
                "period": "Q4",
                "title": "最適化",
                "desc": "効果測定と横展開計画",
            },
        ],
    )

    add_pillars_slide(
        prs,
        title="成功の3本柱",
        subtitle="成果を継続させるための重点領域",
        pillars=[
            {
                "icon_text": "人",
                "title": "People（人）",
                "points": [
                    "現場オーナーを明確化",
                    "役割と権限を定義",
                    "学習機会を定期提供",
                ],
            },
            {
                "icon_text": "流",
                "title": "Process（仕組み）",
                "points": [
                    "標準手順を文書化",
                    "例外対応ルールを整備",
                    "レビュー周期を固定",
                ],
            },
            {
                "icon_text": "基",
                "title": "Platform（基盤）",
                "points": [
                    "データを一元管理",
                    "可視化ダッシュボード",
                    "自動化できる箇所を拡張",
                ],
            },
        ],
    )

    add_matrix_slide(
        prs,
        title="施策ポートフォリオ",
        subtitle="効果と実現容易性のマトリクスで優先順位を決定",
        cells=[
            {
                "title": "Quick Win（すぐやる）",
                "desc": "効果高く、実装も容易。定型報告の自動化やテンプレート整備など、短期間で成果が出る施策を最優先で着手します。",
            },
            {
                "title": "Strategic Bet（戦略投資）",
                "desc": "効果は高いが難易度も高い領域。システム刷新や組織横断連携など、計画的に資源を投下します。",
            },
            {
                "title": "Fill-in（余力で実施）",
                "desc": "容易だが効果は限定的。本丸施策の合間に実施し、現場の小さな不便を解消します。",
            },
            {
                "title": "Avoid / Later（後回し）",
                "desc": "効果が見えにくくコストが高い施策。前提条件が整うまで保留し、定期的に再評価します。",
            },
        ],
    )

    add_takeaways_slide(
        prs,
        title="本日のメッセージ",
        subtitle="意思決定と次アクションにつながる要点",
        items=[
            {
                "number": "01",
                "title": "課題を数値で共有する",
                "desc": "感覚ではなくKPIで現状を揃え、議論の前提をそろえる。",
            },
            {
                "number": "02",
                "title": "小さく試し、早く学ぶ",
                "desc": "大規模導入の前にパイロットで効果とリスクを検証する。",
            },
            {
                "number": "03",
                "title": "人と仕組みを同時に整える",
                "desc": "ツール導入だけでなく、役割・手順・レビューをセットで設計する。",
            },
            {
                "number": "04",
                "title": "横展開可能な形で残す",
                "desc": "成功事例をテンプレート化し、再現性のある改善資産にする。",
            },
        ],
    )

    return prs


def main():
    parser = argparse.ArgumentParser(description="PPT向けインフォグラフィックを生成します")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path("output/sample_infographic.pptx"),
        help="出力先 pptx パス",
    )
    parser.add_argument(
        "-c",
        "--content",
        type=Path,
        default=None,
        help="スライド内容を定義した JSON ファイル（指定時はサンプルではなく JSON から生成）",
    )
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    if args.content:
        prs = build_from_json(args.content)
    else:
        prs = build_sample_deck()
    prs.save(args.output)
    print(f"Generated: {args.output.resolve()}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
