#!/usr/bin/env python3
"""自己紹介インフォグラフィック（1枚）を生成する。"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation

from infographic.self_intro import add_self_intro_image_slide, add_self_intro_slide
from infographic.theme import Theme


DEFAULT_IMAGE_CANDIDATES = [
    Path("assets/self_intro_infographic_ppt.png"),
    Path("assets/self_intro_infographic.png"),
    Path("/opt/cursor/artifacts/assets/self_intro_infographic.png"),
]


def resolve_image(explicit: Path | None) -> Path | None:
    if explicit:
        return explicit if explicit.exists() else None
    for candidate in DEFAULT_IMAGE_CANDIDATES:
        if candidate.exists():
            return candidate
    return None


def resave_with_libreoffice(pptx_path: Path) -> bool:
    """LibreOffice 経由で再保存し、PowerPoint 互換性を高める。"""
    if not shutil.which("soffice"):
        return False
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        cmd = [
            "soffice",
            "--headless",
            "--convert-to",
            "pptx:Impress MS PowerPoint 2007 XML",
            "--outdir",
            str(tmp_dir),
            str(pptx_path),
        ]
        result = subprocess.run(cmd, capture_output=True, text=True)
        converted = tmp_dir / pptx_path.name
        if result.returncode == 0 and converted.exists():
            shutil.copy2(converted, pptx_path)
            return True
    return False


def main():
    parser = argparse.ArgumentParser(description="自己紹介スライドを生成します")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path("output/self_intro.pptx"),
        help="出力 pptx パス",
    )
    parser.add_argument(
        "--image",
        type=Path,
        default=None,
        help="全面配置するインフォグラフィック画像（省略時は assets を自動探索）",
    )
    parser.add_argument(
        "--editable-only",
        action="store_true",
        help="編集可能版のみ生成（画像スライドを含めない）",
    )
    parser.add_argument(
        "--image-only",
        action="store_true",
        help="画像版のみ生成（編集可能スライドを含めない）",
    )
    parser.add_argument(
        "--no-resave",
        action="store_true",
        help="LibreOffice 再保存を行わない",
    )
    args = parser.parse_args()

    prs = Presentation()
    prs.slide_width = Theme.SLIDE_WIDTH
    prs.slide_height = Theme.SLIDE_HEIGHT

    image_path = None if args.editable_only else resolve_image(args.image)
    if image_path:
        add_self_intro_image_slide(prs, image_path)

    if not args.image_only:
        add_self_intro_slide(prs)

    if len(prs.slides) == 0:
        raise SystemExit("生成するスライドがありません")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)

    if not args.no_resave:
        if resave_with_libreoffice(args.output):
            print("Resaved with LibreOffice for PowerPoint compatibility")
        else:
            print("LibreOffice resave skipped/unavailable")

    print(f"Generated: {args.output.resolve()}")
    print(f"Slides: {len(prs.slides)}")
    if image_path:
        print(f"Image slide: {image_path}")
    if not args.image_only:
        print("Editable slide: included")


if __name__ == "__main__":
    main()
