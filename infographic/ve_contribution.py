"""競争吸着 v4：反応一覧・定数・VE出口貢献スライド。"""

from __future__ import annotations

from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .shapes import add_circle, add_rect, add_text_box, header_bar, slide_background
from .slides import add_takeaways_slide, add_title_slide, blank_slide
from .theme import Theme


NAVY = Theme.PRIMARY
GOLD = RGBColor(0xC9, 0xA2, 0x27)
FIT = RGBColor(0xC0, 0x39, 0x2B)
OK = Theme.ACCENT_2
MUTED = Theme.TEXT_MUTED

ASSETS = Path(__file__).resolve().parent.parent / "assets"


def _set_cell(cell, text, *, size=10, bold=False, color=Theme.TEXT, fill=None, align=PP_ALIGN.LEFT):
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = Theme.WHITE
    tf = cell.text_frame
    tf.word_wrap = True
    try:
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.runs[0] if p.runs else p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = Theme.FONT_JP


def _add_table(slide, rows, left, top, width, height, col_w=None):
    n_rows = len(rows)
    n_cols = len(rows[0])
    table = slide.shapes.add_table(n_rows, n_cols, left, top, width, height).table
    if col_w:
        for i, w in enumerate(col_w):
            table.columns[i].width = w
    header = True
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            if header and r_i == 0:
                _set_cell(
                    table.cell(r_i, c_i),
                    str(val),
                    size=10,
                    bold=True,
                    color=Theme.WHITE,
                    fill=NAVY,
                    align=PP_ALIGN.CENTER,
                )
            else:
                highlight = "フィット" in str(val) or str(val).startswith("0.010")
                _set_cell(
                    table.cell(r_i, c_i),
                    str(val),
                    size=10,
                    bold=c_i == 0,
                    color=FIT if highlight else Theme.TEXT,
                    fill=RGBColor(0xFD, 0xF2, 0xE9) if highlight else Theme.WHITE,
                    align=PP_ALIGN.CENTER if c_i > 0 else PP_ALIGN.LEFT,
                )
    return table


def add_image_slide(prs, title, subtitle, image_path: Path, note: str = ""):
    slide = blank_slide(prs)
    header_bar(slide, title, subtitle)
    if image_path.exists():
        top = Inches(1.25)
        height = Inches(5.5 if not note else 5.15)
        slide.shapes.add_picture(str(image_path), Inches(0.45), top, width=Inches(12.4), height=height)
    if note:
        add_text_box(
            slide,
            Inches(0.5),
            Inches(6.5),
            Inches(12.3),
            Inches(0.45),
            note,
            size=Pt(12),
            color=MUTED,
        )
    return slide


def add_glossary_slide(prs):
    slide = blank_slide(prs)
    header_bar(slide, "この資料で使う言葉", "先に略語だけ揃えます。以降は短い名前で呼びます。")
    rows = [
        ("VE1 / α", "ビタミンEのα体（トコフェロールαなど）"),
        ("VE2 / γ", "ビタミンEのγ体"),
        ("FA", "脂肪酸"),
        ("FAEE", "脂肪酸エチルエステル（原料中のエステル）"),
        ("k_f / k_r", "前向きの速さ / 逆向きの速さ"),
        ("Keq", "平衡の定数。逆向きの速さは k_r = k_f / Keq で決まる"),
        ("フィット", "実データに合わせて数値を動かすこと"),
        ("固定", "論文やこれまでの検討の値を、そのまま使うこと"),
        ("出口の山", "塔出口の濃度が、入口より高くなるピーク"),
        ("q_total", "樹脂が持てる吸着量の上限"),
    ]
    for i, (term, meaning) in enumerate(rows):
        col = i % 2
        row = i // 2
        x = Inches(0.45) + Inches(6.4) * col
        y = Inches(1.3) + Inches(1.05) * row
        add_rect(slide, x, y, Inches(6.2), Inches(0.95), Theme.CARD, corner=True)
        add_rect(slide, x, y, Inches(0.14), Inches(0.95), NAVY if col == 0 else Theme.PRIMARY_LIGHT)
        add_text_box(slide, x + Inches(0.35), y + Inches(0.08), Inches(5.7), Inches(0.35), term, size=Pt(14), color=NAVY, bold=True)
        add_text_box(slide, x + Inches(0.35), y + Inches(0.45), Inches(5.7), Inches(0.42), meaning, size=Pt(13), color=Theme.TEXT)
    return slide


def add_reaction_cards_slide(prs):
    slide = blank_slide(prs)
    header_bar(
        slide,
        "7つの反応は、何をしているか",
        "吸着3つ、取り合い3つ（E3は論文にない）、加水分解1つ。逆向きの速さは別に置かず、k_r = k_f / Keq で計算します。",
    )
    groups = [
        (
            "樹脂への吸着（固体と液体のあいだ）",
            Theme.PRIMARY,
            [
                ("A1", "VE1 + S⁺OH⁻ ⇌ S⁺VE1⁻ + OH", "α が樹脂に付く。あとで山になる「溜まり」"),
                ("A2", "VE2 + S⁺OH⁻ ⇌ S⁺VE2⁻ + OH", "γ が樹脂に付く"),
                ("A3", "FA + S⁺OH⁻ ⇌ S⁺FA⁻ + OH", "脂肪酸が樹脂に付く。このあと取り合いのきっかけになる"),
            ],
        ),
        (
            "取り合い（交換）",
            Theme.PRIMARY_LIGHT,
            [
                ("E1", "FA(液) + VE1* ⇌ FA* + VE1(液)", "脂肪酸が α を追い出し、出口へ出す（α の主経路）"),
                ("E2", "FA(液) + VE2* ⇌ FA* + VE2(液)", "脂肪酸が γ を追い出し、出口へ出す（γ の主経路）"),
                ("E3", "VE2(液) + VE1* ⇌ VE2* + VE1(液)", "脂肪酸より先に、γ が α を出す（論文にない反応）"),
            ],
        ),
        (
            "液の中の反応",
            Theme.ACCENT,
            [
                ("H", "FAEE + OH ⇌ FA + ET", "FAEE が分解して脂肪酸が増える（速さ k_hyd だけデータ合わせ）"),
            ],
        ),
    ]
    x = Inches(0.4)
    for title, color, items in groups:
        w = Inches(4.1) if title != "液の中の反応" else Inches(4.15)
        add_rect(slide, x, Inches(1.35), w, Inches(0.42), color, corner=True)
        add_text_box(
            slide,
            x + Inches(0.1),
            Inches(1.38),
            w - Inches(0.2),
            Inches(0.36),
            title,
            size=Pt(13),
            color=Theme.WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        y = Inches(1.9)
        for code, rxn, role in items:
            add_rect(slide, x, y, w, Inches(1.5), Theme.CARD, corner=True)
            add_circle(slide, x + Inches(0.12), y + Inches(0.15), Inches(0.42), color)
            add_text_box(
                slide,
                x + Inches(0.12),
                y + Inches(0.18),
                Inches(0.42),
                Inches(0.38),
                code,
                size=Pt(11),
                color=Theme.WHITE,
                bold=True,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
            )
            add_text_box(
                slide,
                x + Inches(0.62),
                y + Inches(0.12),
                w - Inches(0.75),
                Inches(0.55),
                rxn,
                size=Pt(11),
                color=NAVY,
                bold=True,
            )
            add_text_box(
                slide,
                x + Inches(0.18),
                y + Inches(0.75),
                w - Inches(0.36),
                Inches(0.65),
                role,
                size=Pt(12),
                color=Theme.TEXT,
            )
            y += Inches(1.62)
        x += w + Inches(0.18)

    add_text_box(
        slide,
        Inches(0.45),
        Inches(6.7),
        Inches(12.4),
        Inches(0.55),
        "使っていない向き（前向きの速さ = 0）：液の VE1→樹脂の VE2、液の VE1→樹脂の FA、液の VE2→樹脂の FA。反対向きは、各反応の k_r として自動で入ります。",
        size=Pt(11),
        color=MUTED,
    )
    return slide


def add_policy_slide(prs):
    slide = blank_slide(prs)
    header_bar(
        slide,
        "数値の決め方：動かすものと、そのまま使うもの",
        "データに合わせて動かしているのは、吸着量 q_total と、加水分解の速さ k_hyd だけです。",
    )
    cards = [
        ("データで合わせる", FIT, ["加水分解の速さ k_hyd", "吸着量の上限 q_total", "R003 の k_hyd は範囲の内側", "R004A の k_hyd は上限に当たっている"]),
        ("論文の値を使う", NAVY, ["VE の吸着速さ 0.5838", "脂肪酸の吸着速さ 6.78", "脂肪酸の平衡 412", "単位換算：論文の k × 60"]),
        ("検討して決めた値", Theme.PRIMARY_LIGHT, ["α と γ の差 r=1.7", "脂肪酸の交換速さは論文の2倍", "E3 の速さ 0.6、平衡 1.7", "加水分解の平衡 3.0"]),
        ("逆向きの速さ", Theme.ACCENT_4, ["反対方向の k_f は置かない", "逆向きは k_r = k_f / Keq", "同じ反応を2本にしない", "分配係数 H も動かしていない"]),
    ]
    for i, (title, color, lines) in enumerate(cards):
        x = Inches(0.4) + Inches(3.2) * i
        add_rect(slide, x, Inches(1.4), Inches(3.05), Inches(4.7), Theme.CARD, corner=True)
        add_rect(slide, x, Inches(1.4), Inches(3.05), Inches(0.55), color)
        add_text_box(
            slide,
            x,
            Inches(1.48),
            Inches(3.05),
            Inches(0.4),
            title,
            size=Pt(16),
            color=Theme.WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text_box(
            slide,
            x + Inches(0.18),
            Inches(2.15),
            Inches(2.7),
            Inches(3.7),
            "\n\n".join(f"・{ln}" for ln in lines),
            size=Pt(13),
            color=Theme.TEXT,
        )
    add_text_box(
        slide,
        Inches(0.45),
        Inches(6.3),
        Inches(12.4),
        Inches(0.7),
        "速さ k の単位は cm³ mmol⁻¹ min⁻¹。Keq に単位はありません。最新の合わせ込みは 2026-08-13。R004A の k_hyd は上限です。",
        size=Pt(12),
        color=MUTED,
    )
    return slide


def add_constants_table_slide(prs):
    slide = blank_slide(prs)
    header_bar(slide, "吸着と交換で使っている値", "これらは「一番よい値」ではなく、論文または検討で決めて固定しています。")
    rows = [
        ["定数", "扱い", "論文の値", "いま使っている値", "意味"],
        ["VE の吸着速さ", "固定", "0.5838", "0.5838", "α と γ で分けない"],
        ["脂肪酸の吸着速さ", "固定", "6.78", "6.78", "以前の 0.58 から論文値へ変更"],
        ["VE の吸着平衡 α / γ", "固定", "47.4 / 47.4", "36.3 / 61.8", "差 r=1.7 を採用"],
        ["脂肪酸の吸着平衡", "固定", "412", "412", "論文どおり"],
        ["脂肪酸が VE を出す速さ", "固定", "2.55", "5.1", "論文の2倍。これ以上速くしても効かない"],
        ["その交換の平衡 α / γ", "固定", "324", "422 / 248", "弱い α のほうが脂肪酸に出やすい"],
        ["γ が α を出す速さ", "固定", "なし", "0.6", "0.3〜1.0 の試験の中央"],
        ["その交換の平衡", "固定", "なし", "1.7", "r と同じ"],
    ]
    _add_table(
        slide,
        rows,
        Inches(0.4),
        Inches(1.3),
        Inches(12.5),
        Inches(5.5),
        col_w=[Inches(2.6), Inches(1.5), Inches(2.2), Inches(2.0), Inches(4.2)],
    )
    return slide


def add_fit_values_slide(prs):
    slide = blank_slide(prs)
    header_bar(slide, "データに合わせて動かした値", "反応の速さで動かしているのは、加水分解 k_hyd だけです。")
    kpis = [
        ("0.00346", "k_hyd  R003", "決めた範囲の内側\n以前は 0.00290", Theme.ACCENT_2),
        ("0.010", "k_hyd  R004A", "上限に当たっている\nいちばんよい値ではない", FIT),
        ("0.655", "q_total  R003", "吸着量の上限\n以前は 0.420", NAVY),
        ("1.359", "q_total  R004A", "吸着量の上限\n以前は 0.989", Theme.PRIMARY_LIGHT),
    ]
    for i, (val, label, note, color) in enumerate(kpis):
        x = Inches(0.4) + Inches(3.2) * i
        add_rect(slide, x, Inches(1.45), Inches(3.05), Inches(3.5), Theme.CARD, corner=True)
        add_rect(slide, x, Inches(1.45), Inches(3.05), Inches(0.16), color)
        add_text_box(
            slide, x, Inches(1.85), Inches(3.05), Inches(1.1), val,
            size=Pt(32), color=color, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        add_text_box(
            slide, x + Inches(0.15), Inches(3.05), Inches(2.75), Inches(0.5), label,
            size=Pt(15), color=NAVY, bold=True, align=PP_ALIGN.CENTER,
        )
        add_text_box(
            slide, x + Inches(0.2), Inches(3.6), Inches(2.65), Inches(1.1), note,
            size=Pt(12), color=MUTED, align=PP_ALIGN.CENTER,
        )
    add_rect(slide, Inches(0.4), Inches(5.15), Inches(12.5), Inches(1.75), Theme.CARD, corner=True)
    add_text_box(
        slide,
        Inches(0.65),
        Inches(5.3),
        Inches(12.1),
        Inches(1.45),
        "注意：R004A の k_hyd=0.01 は、出口の誤差を小さくするための上限です。後段の塔でビタミンEを取る目的とは逆向きです。\n"
        "k_hyd を 0 にすると山は平らになります。後段をビタミンE用に残すなら、R003 側の加水分解と脂肪酸吸着が効きます。",
        size=Pt(14),
        color=Theme.TEXT,
    )
    return slide


def add_peak_message_slide(prs):
    slide = blank_slide(prs)
    header_bar(slide, "出口の山は、なぜ入口より高くなるか", "見積もりは R003 の1ロット比較に基づきます。全部の反応を同時に振った解析ではありません。")
    add_rect(slide, Inches(0.45), Inches(1.4), Inches(12.4), Inches(1.55), NAVY, corner=True)
    add_text_box(
        slide,
        Inches(0.7),
        Inches(1.55),
        Inches(12.0),
        Inches(1.25),
        "山の本体は、樹脂に付いていたビタミンEが、脂肪酸に追い出されて液に戻ることです。\n吸着だけだと、出口濃度は入口付近で止まります。",
        size=Pt(20),
        color=Theme.WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    items = [
        ("溜める", "A1 / A2", "ビタミンEを樹脂に付ける。速さだけ変えても、α と γ の山は分かれない"),
        ("増やす", "H + A3", "塔の中で脂肪酸を増やし、追い出しのきっかけをつくる"),
        ("追い出す", "E1 / E2", "樹脂上のビタミンEを液に戻す。入口より高い山の主経路"),
        ("ずらす", "Keq と E3", "出る順番は平衡の差。γ の高さと漏れは E3"),
    ]
    for i, (tag, title, desc) in enumerate(items):
        x = Inches(0.45) + Inches(3.2) * i
        add_rect(slide, x, Inches(3.2), Inches(3.05), Inches(3.5), Theme.CARD, corner=True)
        add_circle(slide, x + Inches(1.1), Inches(3.4), Inches(0.7), Theme.CARD_COLORS[i])
        add_text_box(
            slide, x + Inches(1.1), Inches(3.5), Inches(0.7), Inches(0.5), tag,
            size=Pt(12), color=Theme.WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        add_text_box(
            slide, x + Inches(0.15), Inches(4.25), Inches(2.75), Inches(0.5), title,
            size=Pt(16), color=NAVY, bold=True, align=PP_ALIGN.CENTER,
        )
        add_text_box(
            slide, x + Inches(0.18), Inches(4.85), Inches(2.7), Inches(1.5), desc,
            size=Pt(13), color=Theme.TEXT, align=PP_ALIGN.CENTER,
        )
    return slide


def add_contribution_table_slide(prs):
    slide = blank_slide(prs)
    header_bar(slide, "各反応は、出口の山にどう効くか", "数字はおおよその見積もりです。α の山の主経路は E1、γ の山の主経路は E2 です。")
    rows = [
        ["反応", "役割", "α への効き", "γ への効き", "山の高さ", "出る時刻"],
        ["A1 α の吸着", "α を樹脂に溜める", "必要", "間接", "中", "速さだけでは分かれない"],
        ["A2 γ の吸着", "γ を樹脂に溜める", "間接", "必要", "中", "平衡の差で分かれる"],
        ["A3 脂肪酸の吸着", "取り合いのきっかけ", "間接に大きい", "間接に大きい", "大きい", "中くらい"],
        ["E1 脂肪酸→α", "α を追い出す", "主経路", "副", "大きい", "速さを2倍にしても α は +0.02"],
        ["E2 脂肪酸→γ", "γ を追い出す", "小さい", "主経路", "大きい", "山が少し鋭くなる"],
        ["E3 γ→α", "γ が先に α を出す", "時刻が動く", "高さ・漏れ", "γ が +0.20", "α の山が少し早く出る"],
        ["H 加水分解", "脂肪酸を足す", "間接に大きい", "間接に大きい", "大きい", "R003 なら今の値で足りる"],
    ]
    _add_table(
        slide,
        rows,
        Inches(0.35),
        Inches(1.3),
        Inches(12.6),
        Inches(5.7),
        col_w=[Inches(2.5), Inches(2.3), Inches(1.6), Inches(1.8), Inches(1.8), Inches(2.6)],
    )
    return slide


def add_phase_slide(prs):
    slide = blank_slide(prs)
    header_bar(
        slide,
        "段階を進めると、出口はどう変わったか",
        "R003 の1ロット。実測の山は α 2.80、γ 2.83、間隔 60 分です。",
    )
    rows = [
        ["段階", "変えたもの", "α の山", "γ の山", "間隔(分)", "わかったこと"],
        ["同じ平衡", "α と γ を同じに", "1.77", "1.77", "0", "差がつかない"],
        ["Phase A", "平衡だけ分ける r=1.7", "1.77", "1.95", "72", "出る順番は平衡で決まる。高さはほぼ同じ"],
        ["B1", "E3 を入れる", "1.79", "2.15", "88", "γ の山が高くなった本体"],
        ["B2（いま）", "交換の速さを2倍", "1.81", "2.20", "87", "α の高さは、速さの問題ではない"],
    ]
    _add_table(
        slide,
        rows,
        Inches(0.4),
        Inches(1.3),
        Inches(12.5),
        Inches(3.6),
        col_w=[Inches(2.1), Inches(2.4), Inches(1.6), Inches(1.6), Inches(1.2), Inches(3.6)],
    )
    add_rect(slide, Inches(0.4), Inches(5.15), Inches(12.5), Inches(1.75), Theme.CARD, corner=True)
    add_text_box(
        slide,
        Inches(0.65),
        Inches(5.35),
        Inches(12.1),
        Inches(1.4),
        "実測の山（2.80 / 2.83）には、まだ届いていません。α の高さをどう上げるかは未解決です。\n"
        "交換の速さ、E3、α と γ の差 r を動かしても、α の山はほとんど上がりません。",
        size=Pt(15),
        color=Theme.TEXT,
    )
    return slide


def add_goal_matrix_slide(prs):
    slide = blank_slide(prs)
    header_bar(slide, "何をしたいかで、触る反応が決まる", "次に動かすなら、目的と反応をセットで選んでください。")
    cells = [
        ("入口より高い山にしたい", "効く：加水分解 → 脂肪酸吸着 → 追い出し（E1/E2）", "効かない：原料の FAEE を増やす、交換をさらに速くする"),
        ("α と γ の出る時刻をずらしたい", "効く：吸着と交換の平衡を分ける（Phase A）", "効かない：吸着の速さだけ分ける"),
        ("γ を高く、漏れを減らしたい", "効く：E3 を入れる（B1）", "効かない：α と γ の差 r をさらに下げる"),
        ("α の山を高くしたい", "効くもの：まだ見つかっていない", "効かない：交換の速さ、E3、r"),
    ]
    colors = [NAVY, Theme.PRIMARY_LIGHT, Theme.ACCENT, FIT]
    pos = [
        (Inches(0.4), Inches(1.35)),
        (Inches(6.85), Inches(1.35)),
        (Inches(0.4), Inches(4.25)),
        (Inches(6.85), Inches(4.25)),
    ]
    for (title, good, bad), color, (x, y) in zip(cells, colors, pos):
        add_rect(slide, x, y, Inches(6.05), Inches(2.7), Theme.CARD, corner=True)
        add_rect(slide, x, y, Inches(0.16), Inches(2.7), color)
        add_text_box(slide, x + Inches(0.35), y + Inches(0.18), Inches(5.5), Inches(0.5), title, size=Pt(16), color=color, bold=True)
        add_text_box(slide, x + Inches(0.35), y + Inches(0.8), Inches(5.5), Inches(0.7), good, size=Pt(14), color=Theme.TEXT)
        add_text_box(slide, x + Inches(0.35), y + Inches(1.6), Inches(5.5), Inches(0.8), bad, size=Pt(14), color=MUTED)
    return slide


def build_deck():
    from pptx import Presentation

    prs = Presentation()
    prs.slide_width = Theme.SLIDE_WIDTH
    prs.slide_height = Theme.SLIDE_HEIGHT

    add_title_slide(
        prs,
        title="ビタミンEの出口濃度は\nどの反応で決まっているか",
        subtitle="競争吸着モデル v4　─　反応の一覧、数値の決め方、出口への効き方",
        footer="2026-08-18  |  廣森先生の論文の反応式を実装  |  データ合わせは加水分解の速さと吸着量",
    )
    add_glossary_slide(prs)
    add_takeaways_slide(
        prs,
        title="先に結論",
        subtitle="このあと詳しく見ますが、伝えたいことは4つです。",
        items=[
            {"number": "01", "title": "反応は7つ。データで動かしている速さは加水分解だけ", "desc": "吸着と交換の速さは、論文またはこれまでの検討の値をそのまま使います。吸着量の上限もデータに合わせています。"},
            {"number": "02", "title": "出口の山は、樹脂に付いたビタミンEが追い出されて液に戻ることでできる", "desc": "吸着だけでは入口付近で止まります。山を高くするのは、加水分解 → 脂肪酸吸着 → 追い出し、の順です。"},
            {"number": "03", "title": "出る順番は平衡、γ の高さは E3、α の高さはまだ足りない", "desc": "平衡を分けると時刻がずれます。E3 で γ が高くなります。交換を速くしても α はほとんど上がりません。"},
            {"number": "04", "title": "R004A の加水分解の速さ 0.01 は上限で、いちばんよい値ではありません", "desc": "出口の誤差を小さくするために上限に当たっています。後段でビタミンEを取る目的とは逆です。"},
        ],
    )
    add_image_slide(
        prs,
        "7つの反応（全体像）",
        "吸着3つ、取り合い3つ、液の中の分解1つ",
        ASSETS / "ve_reaction_map.png",
        "E3 は論文にない反応です。速さの定数でデータに合わせているのは、加水分解 k_hyd だけです。",
    )
    add_reaction_cards_slide(prs)
    add_policy_slide(prs)
    add_constants_table_slide(prs)
    add_fit_values_slide(prs)
    add_peak_message_slide(prs)
    add_image_slide(
        prs,
        "山ができる順番",
        "樹脂に溜める → 脂肪酸を増やす → 追い出す → α と γ をずらす",
        ASSETS / "ve_causal_chain.png",
        "入口より高い山の本体は、加水分解 → 脂肪酸吸着 → 追い出しです。E3 は γ を高くし、漏れを減らします。",
    )
    add_contribution_table_slide(prs)
    add_phase_slide(prs)
    add_goal_matrix_slide(prs)
    add_takeaways_slide(
        prs,
        title="まだ足りないこと、次に考えること",
        subtitle="実測の山は α 2.80、γ 2.83。計算はまだそこまで届いていません。",
        items=[
            {"number": "01", "title": "α の山を高くする方法は、まだ見つかっていない", "desc": "交換の速さ、E3、α と γ の差では上がりません。合わせ方の見直しや、α の吸着速さを動かす案があります。"},
            {"number": "02", "title": "R004A の加水分解は、製造のイメージと分けて考える", "desc": "後段をビタミンE用に残すなら、R003 で脂肪酸を足し、R004 の分解は抑えた方がよいです。"},
            {"number": "03", "title": "効き方の数字は、おおよその見積もりです", "desc": "1ロットの比較に基づきます。7つの反応を同時に振った感度解析ではありません。"},
            {"number": "04", "title": "出典", "desc": "Hiromori, Kanuma, Shibasaki-Kitakawa, J. Chem. Eng. Japan, 53(9), 477–484 (2020)."},
        ],
    )
    return prs
