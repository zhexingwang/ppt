"""自己紹介インフォグラフィック（編集可能版）。"""

from __future__ import annotations

from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .shapes import add_circle, add_rect, add_text_box, slide_background
from .theme import Theme


NAVY = RGBColor(0x1F, 0x4E, 0x79)
NAVY_DEEP = RGBColor(0x16, 0x3A, 0x5F)
YELLOW = RGBColor(0xF2, 0xC9, 0x4C)
LIGHT_BG = RGBColor(0xEE, 0xF3, 0xF8)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5A, 0x6A, 0x7A)

ICON_DIR = Path(__file__).resolve().parent.parent / "assets" / "icons"
# PowerPoint互換のため、透明度なしRGB PNGを優先
DEFAULT_ICONS = [
    ICON_DIR / "icon_xian_ppt.png",
    ICON_DIR / "icon_japan_ppt.png",
    ICON_DIR / "icon_lab_ppt.png",
    ICON_DIR / "icon_yokogawa_ppt.png",
]


def _card_with_yellow_bar(slide, left, top, width, height, title: str, subtitle: str):
    add_rect(slide, left, top, width, height, CARD, corner=True)
    add_rect(slide, left, top, Inches(0.12), height, YELLOW)
    add_text_box(
        slide,
        left + Inches(0.22),
        top + Inches(0.12),
        width - Inches(0.35),
        Inches(0.35),
        title,
        size=Pt(13),
        color=NAVY,
        bold=True,
        align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_text_box(
        slide,
        left + Inches(0.22),
        top + Inches(0.45),
        width - Inches(0.35),
        Inches(0.55),
        subtitle,
        size=Pt(10),
        color=MUTED,
        align=PP_ALIGN.LEFT,
    )


def add_self_intro_slide(prs, icon_paths: list[Path] | None = None):
    """王者興さんの自己紹介1枚スライド（図形+イラスト、テキスト編集可能）。"""
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    slide_background(slide, LIGHT_BG)
    icons = list(icon_paths) if icon_paths else list(DEFAULT_ICONS)

    # 上部タイトル帯
    add_rect(slide, 0, 0, Theme.SLIDE_WIDTH, Inches(0.58), NAVY)
    add_text_box(
        slide,
        Inches(0.45),
        Inches(0.06),
        Inches(4),
        Inches(0.45),
        "自己紹介",
        size=Pt(22),
        color=Theme.WHITE,
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    # 氏名カード
    add_rect(slide, Inches(0.45), Inches(0.75), Inches(5.4), Inches(0.95), CARD, corner=True)
    border = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.45),
        Inches(0.75),
        Inches(5.4),
        Inches(0.95),
    )
    border.fill.background()
    border.line.color.rgb = NAVY
    border.line.width = Pt(1.5)
    try:
        border.adjustments[0] = 0.1
    except Exception:
        pass

    add_text_box(
        slide,
        Inches(0.7),
        Inches(0.82),
        Inches(5.0),
        Inches(0.42),
        "王 者興",
        size=Pt(26),
        color=NAVY,
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_text_box(
        slide,
        Inches(0.7),
        Inches(1.25),
        Inches(5.0),
        Inches(0.32),
        "材料科学 × 化学工学 × プロセスシステム工学",
        size=Pt(12),
        color=MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    # 右上補足
    add_text_box(
        slide,
        Inches(6.2),
        Inches(0.9),
        Inches(6.6),
        Inches(0.6),
        "キャリアの歩み ─ 西安から日本、研究、そして協業へ",
        size=Pt(14),
        color=NAVY,
        bold=True,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    # 中央4マイルストーン
    milestones = [
        {
            "num": "01",
            "title": "出身 中国・西安",
            "desc": "かつて「長安」と呼ばれた都市",
        },
        {
            "num": "02",
            "title": "2011 来日",
            "desc": "材料科学で大学卒業後、\n日本語学校で2年学び大学院へ",
        },
        {
            "num": "03",
            "title": "福岡大学大学院",
            "desc": "工学研究科 化学工学・\nプロセスシステム工学 / 野田賢 教授",
        },
        {
            "num": "04",
            "title": "2018- 横河電機",
            "desc": "イノベーションセンター所属\n→ 2025 横河ソリューションサービス",
        },
    ]

    n = 4
    icon_size = Inches(1.35)
    card_w = Inches(2.95)
    card_h = Inches(1.2)
    gap = Inches(0.18)
    total_cards = card_w * n + gap * (n - 1)
    start_x = (Theme.SLIDE_WIDTH - total_cards) / 2
    icon_top = Inches(1.95)
    card_top = Inches(3.55)

    first_center = start_x + card_w / 2
    last_center = start_x + (card_w + gap) * (n - 1) + card_w / 2
    # コネクタ図形は環境によって修復ダイアログの原因になるため、矩形で代替
    line_h = Inches(0.045)
    add_rect(
        slide,
        first_center,
        icon_top + icon_size / 2 - line_h / 2,
        last_center - first_center,
        line_h,
        YELLOW,
    )

    # アイコン間の小さな矢印
    for i in range(n - 1):
        ax = start_x + (card_w + gap) * i + card_w + gap / 2 - Inches(0.12)
        arrow = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_ARROW,
            ax,
            icon_top + icon_size / 2 - Inches(0.1),
            Inches(0.24),
            Inches(0.2),
        )
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = YELLOW
        arrow.line.fill.background()

    for i, m in enumerate(milestones):
        card_x = start_x + (card_w + gap) * i
        ix = card_x + card_w / 2 - icon_size / 2

        # 影代わりの薄い円
        add_circle(
            slide,
            ix + Inches(0.04),
            icon_top + Inches(0.04),
            icon_size,
            RGBColor(0xD0, 0xD8, 0xE0),
        )

        icon_path = icons[i] if i < len(icons) else None
        if icon_path and Path(icon_path).exists():
            slide.shapes.add_picture(
                str(icon_path),
                ix,
                icon_top,
                width=icon_size,
                height=icon_size,
            )
        else:
            add_circle(slide, ix, icon_top, icon_size, NAVY)

        # 番号バッジ（左上）
        badge = Inches(0.32)
        add_circle(slide, ix - Inches(0.02), icon_top - Inches(0.02), badge, YELLOW)
        add_text_box(
            slide,
            ix - Inches(0.02),
            icon_top - Inches(0.02),
            badge,
            badge,
            m["num"],
            size=Pt(10),
            color=NAVY,
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )

        _card_with_yellow_bar(
            slide,
            card_x,
            card_top,
            card_w,
            card_h,
            m["title"],
            m["desc"],
        )

    # 下部タイムラインバー
    bar_top = Inches(5.35)
    bar_h = Inches(1.55)
    add_rect(slide, Inches(0.4), bar_top, Inches(12.5), bar_h, NAVY_DEEP, corner=True)

    add_text_box(
        slide,
        Inches(0.7),
        bar_top + Inches(0.15),
        Inches(4),
        Inches(0.3),
        "年表",
        size=Pt(12),
        color=YELLOW,
        bold=True,
    )

    timeline = [
        ("2011", "中国で大学卒業（材料科学）後に来日\n日本語学校2年 → 大学院進学"),
        ("2018", "博士号取得、横河電機入社\nイノベーションセンター（研究所）"),
        ("2023後半", "PCP様との協業開始\n関連テーマとモデル活用"),
        ("2025", "横河ソリューションサービスへ転籍"),
    ]

    col_w = Inches(2.9)
    for i, (year, text) in enumerate(timeline):
        x = Inches(0.7) + col_w * i
        add_circle(slide, x, bar_top + Inches(0.5), Inches(0.22), YELLOW)
        add_text_box(
            slide,
            x + Inches(0.3),
            bar_top + Inches(0.45),
            col_w - Inches(0.35),
            Inches(0.3),
            year,
            size=Pt(13),
            color=YELLOW,
            bold=True,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        add_text_box(
            slide,
            x + Inches(0.3),
            bar_top + Inches(0.8),
            col_w - Inches(0.35),
            Inches(0.6),
            text,
            size=Pt(10),
            color=Theme.WHITE,
        )

    return slide


def add_self_intro_image_slide(prs, image_path: str | Path):
    """生成済みインフォグラフィック画像を全面配置したスライド。"""
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    path = Path(image_path)
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(
        str(path),
        Inches(0),
        Inches(0),
        width=Theme.SLIDE_WIDTH,
        height=Theme.SLIDE_HEIGHT,
    )
    return slide
