"""各種インフォグラフィック・スライド生成。"""

from __future__ import annotations

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .shapes import (
    add_chevron,
    add_circle,
    add_multiline_box,
    add_rect,
    add_right_arrow,
    add_text_box,
    header_bar,
    slide_background,
)
from .theme import Theme


def blank_slide(prs):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    slide_background(slide)
    return slide


def add_title_slide(
    prs,
    title: str,
    subtitle: str = "",
    footer: str = "",
):
    slide = blank_slide(prs)
    # 左アクセントパネル
    add_rect(slide, 0, 0, Inches(0.35), Theme.SLIDE_HEIGHT, Theme.PRIMARY)
    add_rect(slide, Inches(0.35), 0, Inches(0.12), Theme.SLIDE_HEIGHT, Theme.ACCENT)

    add_text_box(
        slide,
        Inches(1.0),
        Inches(2.4),
        Inches(11.5),
        Inches(1.2),
        title,
        size=Pt(40),
        color=Theme.PRIMARY,
        bold=True,
    )
    if subtitle:
        add_text_box(
            slide,
            Inches(1.0),
            Inches(3.7),
            Inches(11.5),
            Inches(0.7),
            subtitle,
            size=Pt(20),
            color=Theme.TEXT_MUTED,
        )
    if footer:
        add_text_box(
            slide,
            Inches(1.0),
            Inches(6.7),
            Inches(11.5),
            Inches(0.4),
            footer,
            size=Theme.SMALL_SIZE,
            color=Theme.TEXT_MUTED,
        )
    return slide


def add_kpi_slide(
    prs,
    title: str,
    subtitle: str,
    kpis: list[dict],
):
    """
    KPIカード型。
    kpi: {value, label, note?}
    """
    slide = blank_slide(prs)
    header_bar(slide, title, subtitle)

    n = min(len(kpis), 4)
    gap = Inches(0.28)
    left = Theme.MARGIN_X
    usable = Theme.SLIDE_WIDTH - Theme.MARGIN_X * 2 - gap * (n - 1)
    card_w = usable / n
    card_h = Inches(3.6)
    top = Inches(1.4)

    for i, kpi in enumerate(kpis[:n]):
        x = left + (card_w + gap) * i
        color = Theme.CARD_COLORS[i % len(Theme.CARD_COLORS)]

        # カード本体
        add_rect(slide, x, top, card_w, card_h, Theme.CARD, corner=True)
        # 上部カラー帯
        add_rect(slide, x, top, card_w, Inches(0.18), color)

        add_text_box(
            slide,
            x + Inches(0.25),
            top + Inches(0.7),
            card_w - Inches(0.5),
            Inches(1.1),
            str(kpi.get("value", "")),
            size=Pt(42),
            color=color,
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        add_text_box(
            slide,
            x + Inches(0.25),
            top + Inches(2.0),
            card_w - Inches(0.5),
            Inches(0.5),
            kpi.get("label", ""),
            size=Pt(16),
            color=Theme.TEXT,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        if kpi.get("note"):
            add_text_box(
                slide,
                x + Inches(0.25),
                top + Inches(2.55),
                card_w - Inches(0.5),
                Inches(0.7),
                kpi["note"],
                size=Theme.SMALL_SIZE,
                color=Theme.TEXT_MUTED,
                align=PP_ALIGN.CENTER,
            )
    return slide


def add_process_slide(
    prs,
    title: str,
    subtitle: str,
    steps: list[dict],
):
    """
    プロセスフロー（シェブロン）。
    step: {title, desc}
    """
    slide = blank_slide(prs)
    header_bar(slide, title, subtitle)

    n = min(len(steps), 5)
    gap = Inches(0.12)
    left = Theme.MARGIN_X
    usable = Theme.SLIDE_WIDTH - Theme.MARGIN_X * 2 - gap * (n - 1)
    w = usable / n
    h = Inches(1.15)
    top = Inches(1.55)

    for i, step in enumerate(steps[:n]):
        x = left + (w + gap) * i
        color = Theme.CARD_COLORS[i % len(Theme.CARD_COLORS)]
        chevron = add_chevron(slide, x, top, w, h, color)
        # シェブロン内テキスト
        tf = chevron.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        try:
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        except Exception:
            pass
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = f"{i + 1}. {step.get('title', '')}"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = Theme.WHITE
        run.font.name = Theme.FONT_JP

        # 説明カード
        card_top = top + h + Inches(0.35)
        add_rect(
            slide,
            x,
            card_top,
            w,
            Inches(3.4),
            Theme.CARD,
            corner=True,
        )
        add_text_box(
            slide,
            x + Inches(0.18),
            card_top + Inches(0.25),
            w - Inches(0.36),
            Inches(3.0),
            step.get("desc", ""),
            size=Theme.BODY_SIZE,
            color=Theme.TEXT,
            align=PP_ALIGN.LEFT,
        )
    return slide


def add_comparison_slide(
    prs,
    title: str,
    subtitle: str,
    left_panel: dict,
    right_panel: dict,
):
    """
    Before / After 比較。
    panel: {title, items: list[str], color?}
    """
    slide = blank_slide(prs)
    header_bar(slide, title, subtitle)

    panel_w = Inches(5.8)
    panel_h = Inches(5.0)
    top = Inches(1.4)
    left_x = Theme.MARGIN_X
    right_x = Theme.SLIDE_WIDTH - Theme.MARGIN_X - panel_w

    for x, panel, default_color in (
        (left_x, left_panel, Theme.ACCENT_4),
        (right_x, right_panel, Theme.ACCENT_2),
    ):
        color = panel.get("color", default_color)
        add_rect(slide, x, top, panel_w, panel_h, Theme.CARD, corner=True)
        add_rect(slide, x, top, panel_w, Inches(0.7), color)
        add_text_box(
            slide,
            x + Inches(0.3),
            top + Inches(0.12),
            panel_w - Inches(0.6),
            Inches(0.5),
            panel.get("title", ""),
            size=Pt(20),
            color=Theme.WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        items = panel.get("items", [])
        lines = [f"・{item}" for item in items]
        add_multiline_box(
            slide,
            x + Inches(0.4),
            top + Inches(1.0),
            panel_w - Inches(0.8),
            Inches(3.7),
            lines,
            size=Pt(15),
            color=Theme.TEXT,
        )

    # 中央矢印
    add_right_arrow(
        slide,
        Inches(6.25),
        Inches(3.6),
        Inches(0.85),
        Inches(0.45),
        Theme.ACCENT,
    )
    return slide


def add_timeline_slide(
    prs,
    title: str,
    subtitle: str,
    milestones: list[dict],
):
    """
    タイムライン。
    milestone: {period, title, desc}
    """
    slide = blank_slide(prs)
    header_bar(slide, title, subtitle)

    n = min(len(milestones), 5)
    top_line = Inches(3.3)
    left = Inches(0.9)
    right = Inches(12.4)
    # 横線
    add_rect(slide, left, top_line, right - left, Inches(0.08), Theme.PRIMARY_LIGHT)

    span = (right - left) / max(n - 1, 1)
    for i, m in enumerate(milestones[:n]):
        cx = left + span * i
        color = Theme.CARD_COLORS[i % len(Theme.CARD_COLORS)]
        add_circle(slide, cx - Inches(0.18), top_line - Inches(0.14), Inches(0.36), color)

        # 上下交互に配置
        above = i % 2 == 0
        card_w = Inches(2.3)
        card_h = Inches(1.7)
        card_x = cx - card_w / 2
        card_y = (top_line - Inches(2.2)) if above else (top_line + Inches(0.55))

        add_rect(slide, card_x, card_y, card_w, card_h, Theme.CARD, corner=True)
        add_text_box(
            slide,
            card_x + Inches(0.12),
            card_y + Inches(0.12),
            card_w - Inches(0.24),
            Inches(0.35),
            m.get("period", ""),
            size=Pt(12),
            color=color,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text_box(
            slide,
            card_x + Inches(0.12),
            card_y + Inches(0.45),
            card_w - Inches(0.24),
            Inches(0.4),
            m.get("title", ""),
            size=Pt(13),
            color=Theme.TEXT,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text_box(
            slide,
            card_x + Inches(0.12),
            card_y + Inches(0.9),
            card_w - Inches(0.24),
            Inches(0.65),
            m.get("desc", ""),
            size=Pt(11),
            color=Theme.TEXT_MUTED,
            align=PP_ALIGN.CENTER,
        )
    return slide


def add_pillars_slide(
    prs,
    title: str,
    subtitle: str,
    pillars: list[dict],
):
    """
    3本柱 / 4本柱。
    pillar: {icon_text, title, points: list[str]}
    """
    slide = blank_slide(prs)
    header_bar(slide, title, subtitle)

    n = min(len(pillars), 4)
    gap = Inches(0.3)
    usable = Theme.SLIDE_WIDTH - Theme.MARGIN_X * 2 - gap * (n - 1)
    w = usable / n
    top = Inches(1.4)
    h = Inches(5.2)

    for i, pillar in enumerate(pillars[:n]):
        x = Theme.MARGIN_X + (w + gap) * i
        color = Theme.CARD_COLORS[i % len(Theme.CARD_COLORS)]
        add_rect(slide, x, top, w, h, Theme.CARD, corner=True)
        # アイコン円
        circle_size = Inches(0.9)
        add_circle(
            slide,
            x + (w - circle_size) / 2,
            top + Inches(0.35),
            circle_size,
            color,
        )
        add_text_box(
            slide,
            x + (w - circle_size) / 2,
            top + Inches(0.5),
            circle_size,
            Inches(0.6),
            pillar.get("icon_text", str(i + 1)),
            size=Pt(22),
            color=Theme.WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        add_text_box(
            slide,
            x + Inches(0.2),
            top + Inches(1.5),
            w - Inches(0.4),
            Inches(0.6),
            pillar.get("title", ""),
            size=Pt(18),
            color=Theme.PRIMARY,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        points = pillar.get("points", [])
        lines = [f"・{p}" for p in points]
        add_multiline_box(
            slide,
            x + Inches(0.3),
            top + Inches(2.3),
            w - Inches(0.6),
            Inches(2.5),
            lines,
            size=Pt(13),
            color=Theme.TEXT,
        )
    return slide


def add_matrix_slide(
    prs,
    title: str,
    subtitle: str,
    cells: list[dict],
):
    """
    2x2 マトリクス。
    cell: {title, desc, color?}
    """
    slide = blank_slide(prs)
    header_bar(slide, title, subtitle)

    # 必ず4セル想定（足りなければ空埋め）
    padded = (cells + [{"title": "", "desc": ""}] * 4)[:4]
    positions = [
        (Theme.MARGIN_X, Inches(1.35)),
        (Inches(6.85), Inches(1.35)),
        (Theme.MARGIN_X, Inches(4.25)),
        (Inches(6.85), Inches(4.25)),
    ]
    cell_w = Inches(5.95)
    cell_h = Inches(2.6)
    defaults = [
        Theme.PRIMARY,
        Theme.PRIMARY_LIGHT,
        Theme.ACCENT,
        Theme.ACCENT_2,
    ]

    for i, (cell, (x, y)) in enumerate(zip(padded, positions)):
        color = cell.get("color", defaults[i])
        add_rect(slide, x, y, cell_w, cell_h, Theme.CARD, corner=True)
        add_rect(slide, x, y, Inches(0.18), cell_h, color)
        add_text_box(
            slide,
            x + Inches(0.4),
            y + Inches(0.25),
            cell_w - Inches(0.6),
            Inches(0.5),
            cell.get("title", ""),
            size=Pt(18),
            color=color,
            bold=True,
        )
        add_text_box(
            slide,
            x + Inches(0.4),
            y + Inches(0.9),
            cell_w - Inches(0.6),
            Inches(1.4),
            cell.get("desc", ""),
            size=Pt(14),
            color=Theme.TEXT,
        )
    return slide


def add_takeaways_slide(
    prs,
    title: str,
    subtitle: str,
    items: list[dict],
):
    """
    まとめ / Key Takeaways。
    item: {number, title, desc}
    """
    slide = blank_slide(prs)
    header_bar(slide, title, subtitle)

    top = Inches(1.35)
    for i, item in enumerate(items[:5]):
        y = top + Inches(1.1) * i
        color = Theme.CARD_COLORS[i % len(Theme.CARD_COLORS)]
        add_rect(
            slide,
            Theme.MARGIN_X,
            y,
            Theme.SLIDE_WIDTH - Theme.MARGIN_X * 2,
            Inches(0.95),
            Theme.CARD,
            corner=True,
        )
        add_circle(
            slide,
            Theme.MARGIN_X + Inches(0.25),
            y + Inches(0.18),
            Inches(0.6),
            color,
        )
        add_text_box(
            slide,
            Theme.MARGIN_X + Inches(0.25),
            y + Inches(0.28),
            Inches(0.6),
            Inches(0.45),
            str(item.get("number", i + 1)),
            size=Pt(18),
            color=Theme.WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        add_text_box(
            slide,
            Theme.MARGIN_X + Inches(1.1),
            y + Inches(0.12),
            Inches(11.2),
            Inches(0.35),
            item.get("title", ""),
            size=Pt(16),
            color=Theme.PRIMARY,
            bold=True,
        )
        add_text_box(
            slide,
            Theme.MARGIN_X + Inches(1.1),
            y + Inches(0.48),
            Inches(11.2),
            Inches(0.35),
            item.get("desc", ""),
            size=Pt(13),
            color=Theme.TEXT_MUTED,
        )
    return slide
