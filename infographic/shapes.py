"""図形・テキスト描画ヘルパー。"""

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .theme import Theme


def set_fill(shape, color: RGBColor, line: bool = False):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if not line:
        shape.line.fill.background()


def add_rect(slide, left, top, width, height, color: RGBColor, corner=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if corner else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    set_fill(shape, color)
    if corner:
        try:
            shape.adjustments[0] = 0.1
        except Exception:
            pass
    return shape


def add_circle(slide, left, top, size, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    set_fill(shape, color)
    return shape


def add_text_box(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size=Theme.BODY_SIZE,
    color=Theme.TEXT,
    bold=False,
    align=PP_ALIGN.LEFT,
    font_name=Theme.FONT_JP,
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf.vertical_anchor = anchor
    except Exception:
        pass

    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return box


def set_paragraph_text(
    paragraph,
    text: str,
    *,
    size=Theme.BODY_SIZE,
    color=Theme.TEXT,
    bold=False,
    align=PP_ALIGN.LEFT,
    font_name=Theme.FONT_JP,
):
    paragraph.alignment = align
    if paragraph.runs:
        run = paragraph.runs[0]
        run.text = text
        for extra in paragraph.runs[1:]:
            extra.text = ""
    else:
        run = paragraph.add_run()
        run.text = text
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name


def add_multiline_box(
    slide,
    left,
    top,
    width,
    height,
    lines: list[str],
    *,
    size=Theme.BODY_SIZE,
    color=Theme.TEXT,
    bold=False,
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_paragraph_text(
            p,
            line,
            size=size,
            color=color,
            bold=bold,
            align=align,
        )
        p.space_after = Pt(4)
    return box


def add_chevron(slide, left, top, width, height, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, left, top, width, height)
    set_fill(shape, color)
    return shape


def add_right_arrow(slide, left, top, width, height, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    set_fill(shape, color)
    return shape


def slide_background(slide, color=Theme.BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def header_bar(slide, title: str, subtitle: str | None = None):
    """上部アクセントバー + タイトル。"""
    add_rect(slide, 0, 0, Theme.SLIDE_WIDTH, Inches(0.12), Theme.PRIMARY)
    add_text_box(
        slide,
        Theme.MARGIN_X,
        Inches(0.28),
        Inches(12.3),
        Inches(0.55),
        title,
        size=Theme.TITLE_SIZE,
        color=Theme.PRIMARY,
        bold=True,
        align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    if subtitle:
        add_text_box(
            slide,
            Theme.MARGIN_X,
            Inches(0.78),
            Inches(12.3),
            Inches(0.35),
            subtitle,
            size=Theme.SUBTITLE_SIZE,
            color=Theme.TEXT_MUTED,
            bold=False,
        )
